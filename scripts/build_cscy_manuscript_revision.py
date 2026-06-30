from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


BASE = Path(r"C:\Users\M7120\Desktop\System of Eternity\40_SOE_ARXIV_CSCY_REFRAMING_PACKET_2026-06-05")
SOURCE = BASE / "03_SOURCE_REFERENCE_NOT_FOR_UPLOAD" / "SOE_Paper_Manuscript_Draft_v0_7_Codex_source.docx"
OUT = BASE / "01_REVISED_MANUSCRIPT" / "SOE_Paper_Manuscript_Draft_v0_7_csCY_Reframed_2026-06-05.docx"


def add_para_before(paragraph, text, style_name="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.text = text
    new_para.style = style_name
    return new_para


def main():
    doc = Document(str(SOURCE))

    # Front matter: reframe for cs.CY without altering simulation evidence.
    doc.paragraphs[0].text = "Simulation-Bounded Governance Architecture"
    doc.paragraphs[1].text = "Computational Modeling of Institutional Stability in the System of Eternity"
    doc.paragraphs[2].text = (
        "cs.CY arXiv framing revision - simulation evidence unchanged; "
        "companion archive DOI: 10.5281/zenodo.20144235"
    )

    for section in doc.sections:
        for para in section.header.paragraphs:
            if para.text.strip():
                para.text = (
                    "SOE cs.CY manuscript revision - companion archive DOI "
                    "10.5281/zenodo.20144235"
                )
        for para in section.footer.paragraphs:
            if para.text.strip():
                para.text = (
                    "Computational governance framing; simulation evidence unchanged; "
                    "not deployment-ready"
                )

    doc.paragraphs[4].text = (
        "This manuscript presents the System of Eternity (SOE) Governance Architecture as a "
        "computational governance model for studying society-facing institutional stability under "
        "bounded simulation conditions. SOE is represented through theoretical variables T (trust), "
        "D (disturbance), C (cognition), I (identity), G (governance capacity), and S (stability). "
        "The paper consolidates the frozen v0.4 architecture base after Grand Simulation v0.7 "
        "(870 runs; 139,200 step rows; 29 scenarios) and interprets the resulting outputs as "
        "architecture-model evidence for claim-bounded public governance reasoning. In the tested "
        "matrix, federation monitoring uses C07 Psi_extended rather than the C04 S-threshold for "
        "federation Trigger A. Under the locked v0.7 configuration (psi_threshold = 0.30, "
        "stability-heavy weights), Psi_extended produced on-time detections with mean lead 8.175 "
        "steps and zero false positives in audited no-event rows; critically, the detector depends "
        "materially on the stability term. Stability-term ablation reduces on-time TPR to 0.0 and, "
        "at high stress, reverses mean lead to -4.1 steps (CL02A). Hub redundancy is "
        "context-dependent, topology lag can overlap hidden collapse without establishing broad "
        "blind/no-trigger false stability, and resource stress creates floor-duration and "
        "routing-loss gradients without proving recovery under finite-resource pressure. These "
        "results are simulation evidence only: they do not imply deployment readiness, empirical "
        "proxy validity, or governance authority sufficiency. The contribution is a "
        "computationally bounded governance-architecture manuscript with explicit claim-to-output "
        "traceability, allowing simulation-supported public-institutional claims to be separated "
        "from deployment, policy authority, or real-world validation claims."
    )

    doc.paragraphs[6].text = (
        "computers and society; computational governance; governance-system simulation; "
        "institutional modeling; governance architecture; simulation-level evidence; "
        "federation monitoring; stability-dependent detector; topology lag; "
        "resource-constrained recovery; claim-to-output traceability"
    )

    doc.paragraphs[8].text = (
        "SOE is framed as a theoretical computational-governance model rather than a deployment "
        "system. Its purpose is to specify how trust, disturbance, cognition, identity, governance "
        "capacity, and stability interact when a large-scale institutional system faces "
        "topology change, federation monitoring demands, recovery constraints, asynchronous "
        "coordination, and governance stress. The central problem is compositional and "
        "society-facing: a rule that appears safe in an isolated governance module may fail when "
        "coupled to other modules, delayed monitoring, limited resources, or adversarial "
        "governance conditions. This makes the manuscript appropriate for Computers and Society "
        "as a study of computational modeling, claim discipline, and simulation-bounded reasoning "
        "for public institutional architectures."
    )

    doc.paragraphs[9].text = (
        "Contribution statement. Under the Computers and Society framing, this paper contributes "
        "three bounded artifacts: first, a computational governance-architecture model for "
        "studying institutional stability under simulated stress; second, a topology-conditioned "
        "federation-monitoring rule with an explicit stability-dependence caveat; and third, a "
        "claim-to-output traceability framework for separating simulation-supported architecture "
        "claims from public-policy, deployment, or real-world authority claims. The manuscript "
        "does not claim empirical validation of human governance variables or readiness for "
        "real-world authority."
    )

    doc.paragraphs[10].text = (
        "This cs.CY-framed revision preserves the v0.7 simulation evidence, figures, tables, and "
        "claim boundaries while clarifying the manuscript's role as computational governance "
        "research. The prior review found no source-hierarchy contradiction, no deployment "
        "overclaim, and no Grand Simulation reopening trigger. The remaining work is arXiv "
        "submission assembly and external category-fit review. No additional Grand Simulation run "
        "is required unless a hard reproducibility or source-hierarchy defect is found. Core "
        "literature anchors include polycentric and adaptive governance, normal-accident and "
        "interdependent-network theory, distributed-systems fault tolerance, and computational "
        "social-science simulation."
    )

    if doc.tables:
        for row in doc.tables[0].rows:
            label = row.cells[0].text.strip()
            if label == "Readiness status":
                row.cells[1].text = (
                    "cs.CY manuscript framing prepared; simulation evidence consolidated; "
                    "companion archive DOI published; arXiv upload QA and category-fit review "
                    "pending; not deployment-ready."
                )
            elif label == "Patch scope":
                row.cells[1].text = (
                    "v0.7 cs.CY revision reframes title, abstract, keywords, introduction, "
                    "limitations, conclusion, and acknowledgment while preserving Grand "
                    "Simulation v0.7 evidence, figures, tables, and claim boundaries."
                )

    # Add a focused cs.CY positioning paragraph after the conservative framing paragraph.
    after_intro = doc.paragraphs[11]
    p = after_intro.insert_paragraph_before(
        "Category positioning. The paper is framed for cs.CY because it studies how computational "
        "simulation can bound claims about society-facing governance architectures. It is not a "
        "machine-learning paper, a deployed software system, or a public-policy recommendation. "
        "The category fit rests on computational modeling of institutional behavior, public "
        "governance risk, and traceability rules for preventing simulation outputs from being "
        "misread as empirical validation."
    )
    p.style = "Normal"

    # Add a limitation bullet after the first limitation.
    limitations = [i for i, para in enumerate(doc.paragraphs) if para.text.strip() == "6. Limitations"]
    if limitations:
        idx = limitations[0] + 2
        p = doc.paragraphs[idx].insert_paragraph_before(
            "The cs.CY framing identifies the manuscript as computational governance research; "
            "it does not convert SOE into a deployed computing system, public-policy instrument, "
            "or validated socio-technical measurement framework."
        )
        p.style = "List Bullet"

    # Update conclusion to name cs.CY without reopening data claims.
    for para in doc.paragraphs:
        if para.text.startswith("The SOE Governance Architecture v0.4 frozen base remains"):
            para.text = (
                "The SOE Governance Architecture v0.4 frozen base remains a valid foundation for "
                "manuscript drafting. Grand Simulation v0.7 provides sufficient simulation-level "
                "evidence to support conservative claim framing under a Computers and Society "
                "submission path: the manuscript studies computational governance modeling, "
                "institutional-risk simulation, and claim-to-output traceability, not deployment "
                "readiness or empirical governance validation. The next hard gate is final arXiv "
                "submission assembly and external category-fit review; deployment, pilot use, and "
                "real-world measurement remain future work outside this manuscript."
            )
            break

    for para in doc.paragraphs:
        if para.text.startswith("The citation anchors have been inserted."):
            para.text = (
                "The citation anchors have been inserted. For arXiv upload, perform final "
                "source matching and target-style formatting. Each citation should continue "
                "to support a specific claim, method analogy, or limitation rather than "
                "serving as decoration."
            )
            break

    # Add acknowledgment / contribution note before references.
    ref_para = next((p for p in doc.paragraphs if p.text.strip() == "11. References"), None)
    if ref_para is not None:
        add_para_before(ref_para, "11. Acknowledgments and Contribution Note", "Heading 1")
        add_para_before(
            ref_para,
            "QianJun Yu is the sole author responsible for the SOE framework, simulation framing, "
            "manuscript interpretation, and submission decisions. DawLyn Yu is acknowledged for "
            "non-author support and discussion during the broader research and preparation period. "
            "This acknowledgment does not imply responsibility for the technical claims, simulation "
            "outputs, or submission category decisions.",
            "Normal",
        )
        add_para_before(
            ref_para,
            "The author used AI-assisted drafting and review tools for editorial critique, "
            "category-framing review, formatting support, and consistency checks. All claims, "
            "interpretations, references, simulation outputs, and submission decisions remain the "
            "sole responsibility of the author.",
            "Normal",
        )
        ref_para.text = "12. References"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
