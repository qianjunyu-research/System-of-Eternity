from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


BASE_DOC = Path(r"C:\Users\M7120\Desktop\SOE_V4_3_READY_FOR_FINAL_REVIEW_UPLOAD\System_of_Eternity_-_Version_4_3.docx")
FINAL_DOC = Path(r"C:\Users\M7120\Desktop\SOE_V4_3_READY_FOR_FINAL_REVIEW_UPLOAD\System_of_Eternity_-_Version_4_3_final.docx")
FINAL_DOC_SECONDARY = Path(
    r"C:\Users\M7120\Desktop\System of Eternity\SOE Meta-Governance Framework & Civilization\System_of_Eternity_-_Version_4_3_final.docx"
)


HERO_GATE = (
    "Falsifiability gate: Hero activation requires evidence that the system retains "
    "a recovery pathway independent of Hero intervention. If no such pathway exists, "
    "Hero activation is blocked and Trigger B external intervention is the required "
    "response instead. This gate enforces Meta Rule 1: no component of SOE may become "
    "structurally necessary and therefore unchallengeable."
)


CH10_CONSTRAINT = (
    "Emergency and crisis activation procedures in this chapter are subject to "
    "Chapter 3.6 trigger authority constraints without exception. Specifically: "
    "(1) all threshold evaluations require multi-source verification from at least "
    "two non-correlated independent sources; (2) a human review layer is required "
    "before activation; (3) a defined review window must precede irreversible actions; "
    "(4) logged evidence, uncertainty estimate, and source component must be recorded "
    "for each classification decision; (5) emergency powers do not auto-renew — "
    "continuation beyond the initial activation window requires a fresh multi-source "
    "verification cycle. No emergency condition suspends these constraints. Chapter 10 "
    "mechanisms that conflict with Chapter 3.6 are superseded by Chapter 3.6."
)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    new_para.add_run(text)
    return new_para


def main() -> None:
    if not BASE_DOC.exists():
        raise FileNotFoundError(BASE_DOC)

    shutil.copy2(BASE_DOC, FINAL_DOC)
    doc = Document(FINAL_DOC)

    hero_done = HERO_GATE in "\n".join(p.text for p in doc.paragraphs)
    ch10_done = CH10_CONSTRAINT in "\n".join(p.text for p in doc.paragraphs)

    if not hero_done:
        for para in doc.paragraphs:
            text = para.text or ""
            if text.startswith("Trigger Gate — Minimum Operational Definition") and "Conditional fast-trigger exception:" in text:
                updated = text.replace(
                    " Conditional fast-trigger exception:",
                    " " + HERO_GATE + " Conditional fast-trigger exception:",
                    1,
                )
                para.text = updated
                hero_done = True
                break

    if not ch10_done:
        target = None
        for para in doc.paragraphs:
            if (para.text or "").strip() == "If the review does not pass, the emergency state must be terminated.":
                target = para
                break
        if target is None:
            raise ValueError("Chapter 10 activation paragraph not found")
        insert_after(target, CH10_CONSTRAINT, "Normal")
        ch10_done = True

    if not hero_done:
        raise ValueError("Hero trigger gate location not found")
    if not ch10_done:
        raise ValueError("Chapter 10 patch location not found")

    doc.save(FINAL_DOC)
    FINAL_DOC_SECONDARY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FINAL_DOC, FINAL_DOC_SECONDARY)
    print(f"final_doc={FINAL_DOC}")
    print(f"secondary_copy={FINAL_DOC_SECONDARY}")


if __name__ == "__main__":
    main()
