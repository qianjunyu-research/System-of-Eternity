from __future__ import annotations

import argparse
import re
import time
from pathlib import Path
from typing import Dict, Iterable

from deep_translator import GoogleTranslator
from docx import Document


ALPHA_RE = re.compile(r"[A-Za-z]")


def needs_translation(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if not ALPHA_RE.search(stripped):
        return False
    # Skip URL-only paragraphs
    if stripped.startswith("http://") or stripped.startswith("https://"):
        return False
    return True


def split_chunks(text: str, max_len: int = 4000) -> Iterable[str]:
    if len(text) <= max_len:
        yield text
        return
    start = 0
    while start < len(text):
        end = min(start + max_len, len(text))
        if end < len(text):
            split = text.rfind(" ", start, end)
            if split > start:
                end = split
        yield text[start:end]
        start = end


def translate_text(
    text: str,
    translator: GoogleTranslator,
    cache: Dict[str, str],
    retries: int = 4,
) -> str:
    if text in cache:
        return cache[text]

    pieces = []
    for chunk in split_chunks(text):
        translated = None
        for attempt in range(retries):
            try:
                translated = translator.translate(chunk)
                break
            except Exception:
                if attempt == retries - 1:
                    translated = chunk
                else:
                    time.sleep(0.8 * (attempt + 1))
        pieces.append(translated if translated is not None else chunk)
        time.sleep(0.05)

    result = "".join(pieces)
    cache[text] = result
    return result


def translate_document(src: Path, dst: Path) -> None:
    doc = Document(str(src))
    translator = GoogleTranslator(source="auto", target="zh-CN")
    cache: Dict[str, str] = {}

    for p in doc.paragraphs:
        original = p.text
        if needs_translation(original):
            p.text = translate_text(original, translator, cache)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    original = p.text
                    if needs_translation(original):
                        p.text = translate_text(original, translator, cache)

    doc.save(str(dst))


def main() -> None:
    parser = argparse.ArgumentParser(description="Translate DOCX files to Simplified Chinese.")
    parser.add_argument("files", nargs="+", help="Input .docx file paths.")
    parser.add_argument(
        "--suffix",
        default="_zhCN",
        help="Suffix inserted before .docx for translated output files.",
    )
    args = parser.parse_args()

    for raw in args.files:
        src = Path(raw)
        if not src.exists():
            print(f"[MISS] {src}")
            continue
        dst = src.with_name(f"{src.stem}{args.suffix}{src.suffix}")
        print(f"[START] {src}")
        translate_document(src, dst)
        print(f"[DONE]  {dst}")


if __name__ == "__main__":
    main()
