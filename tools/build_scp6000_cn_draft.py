from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUT1 = Path(r"C:\Users\sosoy\Downloads\SCP6000_story_interface_draft_clean.docx")
OUT2 = Path(r"C:\Users\sosoy\OneDrive\桌面\SCP-6000_评论与三元故事接口_首稿_无乱码.docx")


def build() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)

    p = doc.add_paragraph("SCP-6000评论与三元框架升级：故事作为智慧接入接口")
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph("工作定位：结构评论稿（非剧情复述，非哲学抒情）")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    doc.add_heading("摘要", level=1)
    doc.add_paragraph(
        "本文以 SCP-6000 为案例，不讨论“设定是否真实”，而讨论其叙事结构如何触发认知重构。核心结论是："
        "故事并不直接生成结构，但它能够作为认知接口，使抽象结构被体验、被接入、被内化。基于这一观察，"
        "本文提出对三元框架的一项补充：归一、涌现、奇诡负责“结构如何生成”，故事接口负责“结构如何进入主体”。"
        "该补充不替代原三元机制，而是提供从结构到认知主体的连接层。"
    )

    doc.add_heading("1. 评论目标与边界", level=1)
    doc.add_paragraph("本评论的目标不是解释 SCP 设定细节，而是抽取其可迁移的认知结构。")
    doc.add_paragraph("边界如下：")
    doc.add_paragraph("1. 不把故事当作真理证明。", style="List Number")
    doc.add_paragraph("2. 不把故事接口写成结构生成机制。", style="List Number")
    doc.add_paragraph("3. 不破坏三元框架的原有分工。", style="List Number")

    doc.add_heading("2. SCP-6000 的结构性命题", level=1)
    doc.add_paragraph("SCP-6000 给出的关键体验并非“如何阻止终结”，而是“如何重写终结的含义”。")
    doc.add_paragraph("可以压缩为一个结构命题：终结不等于消失，而是叙事层级的转移。")
    doc.add_paragraph("读者在阅读中通常经历如下路径：")
    doc.add_paragraph("灾难理解 -> 对抗期待 -> 对抗失效 -> 结构重释 -> 意义迁移", style="List Bullet")
    doc.add_paragraph("这一变化往往不是靠形式推导完成，而是靠过程体验触发。")

    doc.add_heading("3. 三元框架中的新位置：故事接口层", level=1)
    doc.add_paragraph(
        "在三元框架中，归一、涌现、奇诡负责结构生成问题；但“主体如何真正理解结构”是另一问题。"
    )
    doc.add_paragraph("据此，引入接口层定义：")
    doc.add_paragraph(
        "故事接口层：将抽象结构转化为可体验过程，使主体在经历中完成结构接入。",
        style="List Bullet",
    )
    doc.add_paragraph(
        "简化表达：故事不生产结构，但决定结构能否被理解。",
        style="List Bullet",
    )

    doc.add_heading("4. 与三元机制的关系（防混淆）", level=1)
    doc.add_paragraph("为防止概念漂移，需明确区分：")
    doc.add_paragraph("归一：处理差异并收束结构。", style="List Number")
    doc.add_paragraph("涌现：在规则互动中生成新序。", style="List Number")
    doc.add_paragraph("奇诡：在边界冲突中触发重构。", style="List Number")
    doc.add_paragraph("故事接口：将已有结构接入主体认知。", style="List Number")
    doc.add_paragraph("因此，“故事”在框架中的角色是接入器而非发动机。")

    doc.add_heading("5. 对你体系的实际意义", level=1)
    doc.add_paragraph("这项补充的价值不在于增加概念数量，而在于提升可传播性与可吸收性。")
    doc.add_paragraph("它使你的体系从“懂推导的人才能懂”升级为“非技术读者也能接入”。")
    doc.add_paragraph("在中文平台场景下，这一步尤其关键，因为叙事接入通常优于抽象定义直推。")

    doc.add_heading("6. 最小发布版结论", level=1)
    doc.add_paragraph(
        "SCP-6000 的价值，不在于它提供了一个“答案”，而在于它示范了一个“接口”："
        "当逻辑无法直接完成说服时，故事可以让主体先经历结构，再理解结构。"
    )
    doc.add_paragraph("三元框架因此得到一项最小、但关键的补充：在生成层与结构层之外，增加接口层。")

    doc.add_heading("附：单轮轻审问题（替代 Multi-AI 全量轮）", level=1)
    doc.add_paragraph("仅检查三件事：")
    doc.add_paragraph("1. 是否把故事误写成生成机制。", style="List Number")
    doc.add_paragraph("2. 是否过度哲学化导致不可操作。", style="List Number")
    doc.add_paragraph("3. 是否破坏三元框架的原有分工。", style="List Number")
    return doc


def main() -> None:
    doc = build()
    doc.save(str(OUT1))
    doc.save(str(OUT2))
    print(OUT1)
    print(OUT2)


if __name__ == "__main__":
    main()

