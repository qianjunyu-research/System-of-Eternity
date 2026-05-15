from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ADDENDUM_IN = Path(r"C:\Users\M7120\Downloads\SOE_V4_3_Addendum.docx")
ADDENDUM_OUT = Path(r"C:\Users\M7120\Downloads\SOE_V4_3_Addendum_v2.docx")
V42_IN = Path(
    r"C:\Users\M7120\Desktop\System of Eternity\SOE Meta-Governance Framework & Civilization\System_of_Eternity_-_Version_4_2_REPAIRED_FOR_V4_3_MERGE.docx"
)
V43_OUT = Path(
    r"C:\Users\M7120\Desktop\System of Eternity\SOE Meta-Governance Framework & Civilization\System_of_Eternity_-_Version_4_3.docx"
)


def iter_blocks(doc: DocumentObject):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def block_text(block) -> str:
    if isinstance(block, Paragraph):
        return " ".join((block.text or "").split())
    if isinstance(block, Table):
        vals = []
        for row in block.rows:
            vals.append(" | ".join(" ".join(cell.text.split()) for cell in row.cells))
        return " || ".join(vals)
    return ""


def element(block):
    if isinstance(block, Paragraph):
        return block._p
    return block._tbl


def remove_block(block) -> None:
    el = element(block)
    el.getparent().remove(el)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_before(doc: DocumentObject, ref_para: Paragraph, text: str, style: str | None = None) -> Paragraph:
    para = doc.add_paragraph()
    if style:
        style_obj = None
        for candidate in doc.styles:
            if candidate.name == style or candidate.style_id == style.replace(" ", ""):
                style_obj = candidate
                break
        if style_obj is not None:
            para.style = style_obj
    para.add_run(text)
    ref_para._p.addprevious(para._p)
    return para


def insert_element_before(ref_para: Paragraph, xml_element) -> None:
    ref_para._p.addprevious(xml_element)


def find_paragraph(doc: DocumentObject, predicate) -> Paragraph:
    for para in doc.paragraphs:
        if predicate(" ".join((para.text or "").split())):
            return para
    raise ValueError("Paragraph not found")


def find_block_range(doc: DocumentObject, start_predicate, end_predicate):
    blocks = list(iter_blocks(doc))
    start = None
    end = None
    for idx, block in enumerate(blocks):
        txt = block_text(block)
        if start is None and start_predicate(txt, block):
            start = idx
            continue
        if start is not None and end_predicate(txt, block):
            end = idx
            break
    if start is None or end is None:
        raise ValueError("Block range not found")
    return blocks, start, end


def clone_blocks(blocks):
    return [copy.deepcopy(element(block)) for block in blocks]


def patch_addendum() -> None:
    doc = Document(ADDENDUM_IN)

    # Manual contents list: remove standalone Part E and add folded Section 3.12 under Part A.
    part_a_toc = find_paragraph(doc, lambda t: t == "Part A — Chapter 3 Meta-Governance Layer: Upgraded Edition")
    insert_paragraph_before(doc, doc.paragraphs[15], "    3.12 Long-Term Integrity Monitoring")
    for para in list(doc.paragraphs):
        if " ".join((para.text or "").split()) == "Part E — Drift Layer: Cross-Cutting Architecture Specification":
            # Keep the actual Part E heading for now; this first hit is the contents item.
            if para is not None and para._p is not part_a_toc._p:
                remove_block(para)
                break

    # Source wording softening.
    for para in doc.paragraphs:
        if "confirmed this operationally" in para.text:
            set_paragraph_text(
                para,
                para.text.replace(
                    "confirmed this operationally",
                    "simulation-supported within the tested model",
                ),
            )
        if "simulation-validated exclusion zones" in para.text:
            set_paragraph_text(
                para,
                para.text.replace(
                    "simulation-validated exclusion zones",
                    "simulation-supported ranges within the tested model",
                ),
            )

    # Fold Drift Layer into Meta-Governance as Section 3.12.
    drift_table = doc.tables[15]
    drift_table_xml = copy.deepcopy(drift_table._tbl)

    relationship_heading = find_paragraph(doc, lambda t: t == "3.12 Relationship to Other SOE Layers")
    set_paragraph_text(relationship_heading, "3.13 Relationship to Other SOE Layers")
    claim_heading = find_paragraph(doc, lambda t: t == "3.13 Grand Simulation v0.7 Claim Locks")
    set_paragraph_text(claim_heading, "3.14 Grand Simulation v0.7 Claim Locks")
    conclusion_heading = find_paragraph(doc, lambda t: t == "3.14 Chapter Conclusion")
    set_paragraph_text(conclusion_heading, "3.15 Chapter Conclusion")

    insert_paragraph_before(doc, relationship_heading, "")
    insert_paragraph_before(doc, relationship_heading, "3.12 Long-Term Integrity Monitoring", "Heading 2")
    insert_paragraph_before(
        doc,
        relationship_heading,
        "The Drift Layer is not a separate governance component. Drift detection is a Meta-Governance sub-function. It detects slow deviation from SOE's foundational commitments that acute triggers miss. The following drift types are monitored under Meta-Governance authority with the same five-function separation required in Section 3.7.",
    )
    insert_element_before(relationship_heading, drift_table_xml)
    insert_paragraph_before(doc, relationship_heading, "Anti-Capture Rules for Drift Monitoring", "Heading 3")
    anti_capture_items = [
        "The Drift Layer must not become a hidden veto authority. Anti-capture rules:",
        "No single actor controls drift classification.",
        "Drift findings are reviewable by any network participant.",
        "Dismissed severe findings require written rationale, not silence.",
        "Drift reviewers rotate or have an independent review path.",
        "Major pause, rollback, or recognition denial triggered by Drift Layer findings requires independent review.",
    ]
    for idx, item in enumerate(anti_capture_items):
        style = None if idx == 0 else "List Paragraph"
        insert_paragraph_before(doc, relationship_heading, item, style)
    insert_paragraph_before(
        doc,
        relationship_heading,
        "The five-function authority separation defined in Section 3.7 applies to drift monitoring. No single actor may perform sensing, classification, routing, response, and audit for the same drift finding. Drift response actions that involve major pause, rollback, or recognition denial require independent review under the same rules as Trigger C.",
    )
    insert_paragraph_before(doc, relationship_heading, "Open Items Before Simulation Readiness", "Heading 3")
    open_items = [
        "The Drift Layer requires the following before it can be simulation-ready:",
        "Quantitative drift indicators — what measurable signals indicate each drift type?",
        "Detection interval specification — how frequently is each drift type assessed?",
        "Severity classification — what constitutes a minor drift finding vs a Trigger C-level drift finding?",
        "Drift response protocols — what are the correct responses to each drift type?",
        "Drift Layer simulation specification — how does drift interact with the grand simulation model?",
    ]
    for idx, item in enumerate(open_items):
        style = None if idx == 0 else "List Paragraph"
        insert_paragraph_before(doc, relationship_heading, item, style)
    insert_paragraph_before(doc, relationship_heading, "")

    # Internal references to the former Part E.
    for para in doc.paragraphs:
        if "see Part E" in para.text:
            set_paragraph_text(para, para.text.replace("see Part E", "see Section 3.12"))
        if "Drift Layer (Part E)" in para.text:
            set_paragraph_text(
                para,
                para.text.replace("Drift Layer (Part E)", "Meta-Governance drift monitoring (Section 3.12)"),
            )

    # Delete standalone Part E block after folding.
    blocks, start, end = find_block_range(
        doc,
        lambda t, b: isinstance(b, Paragraph) and t == "Part E — Drift Layer: Cross-Cutting Architecture Specification",
        lambda t, b: isinstance(b, Paragraph) and t == "Appendix — Change Summary and V4.3 Merge Instructions",
    )
    for block in blocks[start:end]:
        remove_block(block)

    # C09-G05.
    for para in doc.paragraphs:
        if "Star-adjacent structures — where one participant controls all connections — must be redesigned before scaling." in para.text:
            set_paragraph_text(
                para,
                para.text.replace(
                    "Star-adjacent structures — where one participant controls all connections — must be redesigned before scaling.",
                    "Star-adjacent structures — where one participant controls all connections — must be redesigned before recognition is granted. Hub redundancy is not a substitute for topology redesign. A node may use hub redundancy only as a documented transitional configuration with a mandatory reclassification timeline. Mesh or federation topology is the required destination. Recognition at Stage 2 requires that the final topology is not star-adjacent.",
                ),
            )

    # C09-G06 third-party audit requirement.
    c09_g07 = find_paragraph(doc, lambda t: t == "C09-G07 — Rollback Readiness")
    insert_paragraph_before(
        doc,
        c09_g07,
        "Gate satisfaction for C09-G06 requires verification by at least one independent party not involved in the node's formation. Self-declaration alone is insufficient. The independent verifier must produce a written attestation that no single actor controls all recognition gates. This attestation is logged in the node's recognition record.",
    )
    insert_paragraph_before(doc, c09_g07, "")

    # Source-fidelity caveat in claim locks opening.
    for para in doc.paragraphs:
        if para.text.startswith("The following claim statuses are locked from Grand Simulation v0.7."):
            set_paragraph_text(
                para,
                "The following claim statuses are locked from Grand Simulation v0.7. Grand Simulation v0.7 tested an operationalized C00-C08 model, not the full 201-page framework text. All claims in this section are scoped to the tested model. They do not constitute validation of the full SOE framework. They are simulation-level evidence only. They do not establish deployment readiness.",
            )
            break

    # Federation row in topology detection table.
    topology_table = doc.tables[2]
    if not any("Federation" in row.cells[0].text for row in topology_table.rows):
        row = topology_table.add_row()
        row.cells[0].text = "Federation"
        row.cells[1].text = (
            "Two-tier aggregation. S_network does not accumulate sufficient signal to breach C04 S-threshold by design. "
            "C07 Psi_extended is the mandatory detection pathway."
        )
        row.cells[2].text = (
            "C07 Psi_extended only. C04 S-threshold does not activate Trigger A for federation topology. "
            "psi_lambda = 0.10 locked. psi_threshold = 0.30 locked (v0.7)."
        )

    # Node v0.1 execution period correction and versioned note.
    meta_table = None
    for table in doc.tables:
        for row in table.rows:
            if (
                len(row.cells) >= 2
                and row.cells[0].text.strip() == "Document"
                and "SOE_Node_v0_1_Completion_Report" in row.cells[1].text
            ):
                meta_table = table
                break
        if meta_table is not None:
            break
    if meta_table is None:
        raise ValueError("Node v0.1 metadata table not found")
    # Clean up any prior misplaced correction note from recognition-stage tables.
    for table in doc.tables:
        for row in list(table.rows):
            if row.cells and row.cells[0].text.strip() == "Versioned correction note":
                table._tbl.remove(row._tr)
    for row in meta_table.rows:
        if row.cells[0].text.strip() == "Execution Period":
            row.cells[1].text = "May 6-12, 2026 (7 days)"
    if not any("Versioned correction note" in row.cells[0].text for row in meta_table.rows):
        row = meta_table.add_row()
        row.cells[0].text = "Versioned correction note"
        row.cells[1].text = (
            "v0.1 typo correction: execution period corrected from 'May 6-8212, 2026' to "
            "'May 6-12, 2026' during multi-AI review patch. Frozen v0.1 record otherwise unchanged."
        )

    # Appendix A.1 merge instruction row and V4.3 version block wording.
    change_table = doc.tables[16]
    for row in change_table.rows:
        if row.cells[0].text.strip() == "Drift Layer specification":
            row.cells[0].text = "Drift Layer folded into Meta-Governance Chapter 3"
            row.cells[1].text = "Folded into Meta-Governance Chapter 3 as Section 3.12. Not a standalone component."
            row.cells[2].text = "No standalone chapter or appendix insertion. Retain architecture-specified-only status."

    version_table = doc.tables[17]
    for row in version_table.rows:
        key = row.cells[0].text.strip()
        if key == "Primary additions":
            row.cells[1].text = (
                "C09 Node Formation, Chapter 3 Meta-Governance upgrade, Node v0.1 execution record, "
                "Node Measurement Protocol requirements, Drift Layer folded into Chapter 3 as Section 3.12"
            )
        if key == "Open items":
            row.cells[1].text = (
                "Node Measurement Protocol (not written), C09 simulation spec (not run), drift monitoring simulation spec "
                "(not run), Dynamic Architecture chapters not yet upgraded"
            )

    doc.save(ADDENDUM_OUT)


def extract_section_blocks(doc: DocumentObject, start_text: str, end_texts: tuple[str, ...]):
    blocks = list(iter_blocks(doc))
    start = None
    end = None
    for idx, block in enumerate(blocks):
        txt = block_text(block)
        is_heading = (
            isinstance(block, Paragraph)
            and block.style is not None
            and block.style.name.startswith("Heading")
        )
        if start is None and is_heading and (txt == start_text or txt.startswith(start_text)):
            start = idx
            continue
        if start is not None and is_heading and any(txt == end or txt.startswith(end) for end in end_texts):
            end = idx
            break
    if start is None or end is None:
        raise ValueError(f"Could not extract section {start_text!r}")
    return blocks[start:end]


def patch_first_heading(xml_blocks, replacement: str):
    tmp_doc = Document()
    # Direct XML text replacement is simpler and preserves the copied heading style.
    for block in xml_blocks:
        for text_el in block.iter(qn("w:t")):
            if text_el.text:
                text_el.text = replacement
                return


def merge_v43() -> None:
    addendum = Document(ADDENDUM_OUT)
    base = Document(V42_IN)

    # Front matter version block and change log.
    front_ref = base.paragraphs[3]
    insert_paragraph_before(base, front_ref, "Version 4.3 | May 2026")
    insert_paragraph_before(base, front_ref, "Supersedes SOE Version 4.2 (Zenodo: 10.5281/zenodo.20047391)")
    insert_paragraph_before(base, front_ref, "Simulation base: Grand Simulation v0.7 (Zenodo: 10.5281/zenodo.20144235)")
    insert_paragraph_before(base, front_ref, "Readiness: architecture-review-ready; not pilot-ready or deployment-ready.")
    insert_paragraph_before(base, front_ref, "")
    insert_paragraph_before(base, front_ref, "V4.3 Change Log", "Heading 1")
    for item in [
        "Chapter 3 Meta-Governance upgraded with Grand Simulation v0.7 bounded claim locks.",
        "C09 Node Formation and Recognition Layer added as architecture-specified only.",
        "Node v0.1 execution record added with container-test boundary.",
        "Node Measurement Protocol requirements added for future Node v0.2 work.",
        "Drift monitoring folded into Meta-Governance Chapter 3 as Section 3.12.",
    ]:
        insert_paragraph_before(base, front_ref, item, "List Paragraph")
    insert_paragraph_before(base, front_ref, "")

    # Manual TOC/dynamic architecture front matter updates.
    try:
        next_para = find_paragraph(base, lambda t: t == "Appendices")
    except Exception:
        next_para = find_paragraph(base, lambda t: t == "Appendix A — Technical Infrastructure and Stress Tests")
    insert_paragraph_before(
        base,
        next_para,
        "C09 — Node Formation and Recognition Layer (architecture-specified only)",
        "Heading 2",
    )
    insert_paragraph_before(base, next_para, "SOE Execution Record — Node v0.1 (container evidence only)", "Heading 2")
    insert_paragraph_before(base, next_para, "Node Measurement Protocol Requirements", "Heading 2")

    # Replace Dynamic Architecture Chapter 3.
    add_part_a = clone_blocks(
        extract_section_blocks(
            addendum,
            "Part A — Chapter 3 Meta-Governance Layer: Upgraded Edition",
            ("Part B — C09: Node Formation and Recognition Layer",),
        )
    )
    patch_first_heading(add_part_a, "Chapter 3 — Meta-Governance Layer: Upgraded Edition")
    base_blocks, start, end = find_block_range(
        base,
        lambda t, b: isinstance(b, Paragraph) and t == "Chapter 3 — Meta Governance Layer",
        lambda t, b: isinstance(b, Paragraph) and t == "Chapter 4 — Coordination Layer",
    )
    ref_block = base_blocks[end]
    for xml in add_part_a:
        element(ref_block).addprevious(xml)
    for block in base_blocks[start:end]:
        remove_block(block)

    # Insert Part B-D before Appendix A.
    add_b = clone_blocks(
        extract_section_blocks(
            addendum,
            "Part B — C09: Node Formation and Recognition Layer",
            ("Part C — SOE Execution Record: Node v0.1",),
        )
    )
    add_c = clone_blocks(
        extract_section_blocks(
            addendum,
            "Part C — SOE Execution Record: Node v0.1",
            ("Part D — Node Measurement Protocol Requirements",),
        )
    )
    add_d = clone_blocks(
        extract_section_blocks(
            addendum,
            "Part D — Node Measurement Protocol Requirements",
            ("Appendix — Change Summary and V4.3 Merge Instructions",),
        )
    )
    patch_first_heading(add_b, "C09 — Node Formation and Recognition Layer")
    patch_first_heading(add_c, "SOE Execution Record — Node v0.1")
    patch_first_heading(add_d, "Node Measurement Protocol Requirements")

    appendix_a = find_paragraph(base, lambda t: t == "Appendix A — Technical Infrastructure and Stress Tests")
    for xml in add_b + add_c + add_d:
        appendix_a._p.addprevious(xml)

    base.save(V43_OUT)


def main() -> None:
    if not ADDENDUM_IN.exists():
        raise FileNotFoundError(ADDENDUM_IN)
    if not V42_IN.exists():
        raise FileNotFoundError(V42_IN)
    patch_addendum()
    merge_v43()
    print(f"patched_addendum={ADDENDUM_OUT}")
    print(f"merged_v43={V43_OUT}")


if __name__ == "__main__":
    main()
