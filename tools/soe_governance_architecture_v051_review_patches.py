from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SRC = Path(
    r"C:\Users\M7120\Desktop\SOE_Governance_Architecture_Post_V4_3_ALIGNMENT_UPLOAD_TO_DRIVE\SOE_Architecture_v2_Integration_Snapshot_v0_5_1_POST_V4_3_ALIGNMENT_CANDIDATE.docx"
)
OUT = Path(
    r"C:\Users\M7120\Documents\New project 7\00_ACTIVE_HANDOFFS\SOE_Upgrade_Waiting_Room\SOE_Architecture_v2_Integration_Snapshot_v0_5_1_FREEZE_CANDIDATE.docx"
)
DESKTOP_FULL = Path(
    r"C:\Users\M7120\Desktop\System of Eternity\SOE Governance Architecture V2\SOE Governance Architecture v0.5 Full Snapshot Upgrade\SOE_Architecture_v2_Integration_Snapshot_v0_5_1_FREEZE_CANDIDATE.docx"
)
UPLOAD = Path(r"C:\Users\M7120\Desktop\SOE_Governance_Architecture_Post_V4_3_ALIGNMENT_UPLOAD_TO_DRIVE")


EVIDENCE_GUARD = (
    "Evidence Terminology Guard: Historical terms such as \"validated,\" \"locked,\" \"confirmed,\" "
    "or \"empirical constraint\" inside the C00-C08 integration snapshot refer to simulation-supported "
    "architecture constraints within the tested SOE simulation lineage. They do not imply empirical "
    "validation, pilot readiness, deployment readiness, real-world safety, or validated measurement proxies."
)


HERO_ANTI_CIRCULARITY = (
    "H(t)/Hero intervention may support recovery only when an independently observable non-Hero recovery "
    "pathway exists; Hero activity itself cannot be counted as proof that such a pathway exists."
)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    new_para.add_run(text)
    return new_para


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    # Cover/front-matter traceability.
    if len(doc.paragraphs) > 5:
        doc.paragraphs[1].text = "Governance Architecture v0.5.1 — Post-V4.3 Alignment Freeze Candidate"
        doc.paragraphs[2].text = "Full Integration Snapshot Line — C00-C08 Preserved; v0.5.1 Alignment Applied"
        doc.paragraphs[5].text = "2026-05-15"
        doc.paragraphs[7].text = (
            "Integration Checkpoint: PASSED — C00-C08 preserved; post-V4.3 alignment patches applied"
        )
    if not any("Evidence Terminology Guard:" in p.text for p in doc.paragraphs):
        anchor = doc.paragraphs[10]
        current = insert_after(anchor, "")
        for text in [
            "v0.5.1 Change Summary",
            "C09-G05 strengthened: star-adjacent topology blocks recognition, not merely scaling.",
            "C09-G06 strengthened: independent written anti-capture attestation required.",
            "Drift reframed as C07 Long-Term Integrity Monitoring, not a standalone governance component.",
            "H(t)/Hero falsifiability and anti-circularity safeguards added.",
            "V4.3 Chapter 10 crisis text was not inserted into the Governance Architecture snapshot.",
            EVIDENCE_GUARD,
            "Visual QA Note: Final visual QA should be performed manually in Word/PDF export because automated DOCX-to-PNG render verification was unavailable in the verification environment.",
            "",
        ]:
            current = insert_after(current, text)

    # Manifest traceability for the alignment material.
    manifest_items = {
        "C09 — Node Formation and Recognition Layer (v0.5.1 proposed extension)",
        "C07 Long-Term Integrity Monitoring / Drift Requirements (v0.5.1 proposed sub-function)",
        "H(t)/Hero falsifiability safeguard (post-V4.3 alignment)",
    }
    existing = {p.text for p in doc.paragraphs[:40]}
    if not manifest_items.issubset(existing):
        anchor = None
        for p in doc.paragraphs[:40]:
            if p.text == "C08 — Coordination Layer v17 Addendum":
                anchor = p
                break
        if anchor is not None:
            current = anchor
            for item in manifest_items:
                if item not in existing:
                    current = insert_after(current, item, "List Bullet")

    # Hero anti-circularity and C09 verifier independence.
    for para in doc.paragraphs:
        if para.text.startswith("Post-V4.3 falsifiability constraint") and HERO_ANTI_CIRCULARITY not in para.text:
            para.text = para.text + " " + HERO_ANTI_CIRCULARITY
        if para.text.startswith("C09-G06 Anti-capture:") and "must not be a founder" not in para.text:
            para.text = (
                para.text
                + " For C09-G06, an independent verifier must not be a founder, funder, operator, direct participant, "
                "dependent contractor, recognition beneficiary, or subordinate of the candidate node. Any conflict of "
                "interest must be disclosed in the recognition record."
            )
        if para.text.startswith("C09-G05 Topology fitness:") and "Missed reclassification" not in para.text:
            para.text = (
                para.text
                + " A transitional hub-redundant configuration cannot receive Stage 2 recognition and must pause "
                "recognition or scaling review if the reclassification timeline is missed."
            )
        if para.text.startswith("C07 long-term integrity monitoring detects") and "Quantitative drift thresholds" not in para.text:
            insert_after(
                para,
                "Quantitative drift thresholds remain open simulation and measurement requirements. v0.5.1 does not freeze numeric drift thresholds without a reviewed measurement protocol and simulation matrix.",
            )
        if para.text == "Before Node v0.2 or any pilot-facing claim, SOE requires a standalone Node Measurement Protocol.":
            insert_after(
                para,
                "Node v0.2 is blocked until the Node Measurement Protocol is reviewed, adopted, and multi-AI validated. Node v0.1 remains a bounded container test only and implies no dynamic validation.",
            )

    # Tables with old H(t) availability wording.
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if not cells:
                continue
            row_text = " || ".join(c.text for c in cells)
            if "Only H(t) and R_base can break this" in row_text:
                for cell in cells:
                    cell.text = cell.text.replace(
                        "Only H(t) and R_base can break this.",
                        "R_base is the baseline active mechanism. H(t) is blocked unless an independent non-Hero recovery pathway exists; otherwise Trigger B external intervention is required.",
                    )
            if cells[0].text == "GM-A: Trust-governance deadlock":
                cells[1].text = (
                    "T → 0 and G(t) → 0 simultaneously. γT term collapses. G(t) degrades. Governance cannot function. "
                    "R_base is the only remaining baseline active mechanism. H(t) is blocked unless an independent "
                    "non-Hero recovery pathway is confirmed; otherwise Trigger B external intervention is required."
                )
                if len(cells) > 2:
                    cells[2].text = (
                        "CONFIRMED — architecture. R_base is the designed baseline response; H(t) is conditional on "
                        "the post-V4.3 falsifiability gate and cannot be treated as structurally necessary."
                    )
            if len(cells) >= 2 and cells[0].text == "H(t) requires independent recovery pathway":
                cells[1].text = (
                    cells[1].text
                    + " "
                    + HERO_ANTI_CIRCULARITY
                    if HERO_ANTI_CIRCULARITY not in cells[1].text
                    else cells[1].text
                )
            if len(cells) >= 2 and cells[0].text == "H(t) activation criteria":
                cells[1].text = (
                    "When does H(t) activate? What triggers it? What are the time limits? What is the trigger "
                    "authority structure? Criteria must include the post-V4.3 falsifiability gate and anti-circularity "
                    "rule: Hero activity itself cannot count as proof of a non-Hero recovery pathway."
                )
            if len(cells) >= 2 and cells[0].text == "Response package":
                if "H(t) may be activated (see C06)" in cells[1].text:
                    cells[1].text = cells[1].text.replace(
                        "H(t) may be activated (see C06).",
                        "H(t) may be activated only if the post-V4.3 falsifiability gate is satisfied (see C06).",
                    )
                if "H(t) activated if not already." in cells[1].text:
                    cells[1].text = cells[1].text.replace(
                        "H(t) activated if not already.",
                        "H(t) may activate only if the post-V4.3 falsifiability gate is satisfied; if no independent non-Hero recovery pathway exists, H(t) is blocked and Trigger B external intervention remains the required response.",
                    )

    doc.save(OUT)
    DESKTOP_FULL.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, DESKTOP_FULL)
    UPLOAD.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, UPLOAD / OUT.name)
    print(f"freeze_candidate={OUT}")
    print(f"desktop_copy={DESKTOP_FULL}")
    print(f"upload_copy={UPLOAD / OUT.name}")


if __name__ == "__main__":
    main()
