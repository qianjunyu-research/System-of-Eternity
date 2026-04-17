from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


SRC = Path(r"C:\Users\sosoy\OneDrive\桌面\RAW STORY.docx")
OUT = SRC.with_name("RAW STORY_formatted.docx")


def looks_like_heading(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if len(t) > 28:
        return False
    if any(ch in t for ch in "。！？.!?；;，,:："):
        return False
    return True


def main() -> None:
    doc = Document(str(SRC))

    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.05)
        sec.right_margin = Inches(1.05)

    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    # First non-empty paragraph -> title style
    first_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            first_idx = i
            break

    if first_idx is not None:
        p = doc.paragraphs[first_idx]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(14)
        for r in p.runs:
            r.font.name = "Microsoft YaHei"
            r.font.size = Pt(18)
            r.bold = True

    # Paragraph pass
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if i == first_idx:
            continue

        style_name = p.style.name.lower() if p.style and p.style.name else ""
        is_heading = ("heading" in style_name) or looks_like_heading(t)

        if is_heading:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            for r in p.runs:
                r.font.name = "Microsoft YaHei"
                r.font.size = Pt(13)
                r.bold = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.28)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            for r in p.runs:
                r.font.name = "Microsoft YaHei"
                r.font.size = Pt(11)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if not p.text.strip():
                        continue
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.first_line_indent = Inches(0.2)
                    p.paragraph_format.line_spacing = 1.4
                    for r in p.runs:
                        r.font.name = "Microsoft YaHei"
                        r.font.size = Pt(10.5)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()

