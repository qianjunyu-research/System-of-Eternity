from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


TARGETS = [
    Path(r"C:\Users\sosoy\OneDrive\桌面\System of Eternity\_SOE_Falsification_Log_v1_格式化版.docx"),
    Path(r"C:\Users\sosoy\OneDrive\桌面\System of Eternity\_SOE_Falsification_Log_v1_格式化版_zhCN.docx"),
    Path(r"C:\Users\sosoy\OneDrive\桌面\System of Eternity\SOE_Model_v1.3docx.docx"),
    Path(r"C:\Users\sosoy\OneDrive\桌面\System of Eternity\SOE_Model_v1.3docx_zhCN.docx"),
]


def is_zh_file(path: Path) -> bool:
    name = path.name.lower()
    return ("zhcn" in name) or any("\u4e00" <= ch <= "\u9fff" for ch in path.name)


def looks_like_heading(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if len(t) > 34:
        return False
    if any(p in t for p in "。！？.!?；;，,"):
        return False
    return True


def set_run_font(run, font_name: str, size: float = 11.0, bold: bool | None = None):
    run.font.name = font_name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def apply_para_body_style(para, font_name: str):
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.first_line_indent = Inches(0.28)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.5
    for run in para.runs:
        set_run_font(run, font_name, 11)


def apply_para_heading_style(para, font_name: str):
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.first_line_indent = Inches(0)
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.2
    for run in para.runs:
        set_run_font(run, font_name, 13, bold=True)


def beautify_doc(src: Path) -> Path:
    doc = Document(str(src))

    is_zh = is_zh_file(src)
    body_font = "Microsoft YaHei" if is_zh else "Calibri"
    title_font = "Microsoft YaHei" if is_zh else "Calibri"

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.05)
        section.right_margin = Inches(1.05)

    # Base normal style
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    # Title pass
    first_non_empty = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            first_non_empty = i
            break

    if first_non_empty is not None:
        p = doc.paragraphs[first_non_empty]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.space_before = Pt(2)
        for run in p.runs:
            set_run_font(run, title_font, 18, bold=True)

    # Body pass
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if i == first_non_empty:
            continue
        style_name = p.style.name.lower() if p.style and p.style.name else ""
        if "heading" in style_name or looks_like_heading(t):
            apply_para_heading_style(p, title_font)
        else:
            apply_para_body_style(p, body_font)

    # Tables pass
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        apply_para_body_style(p, body_font)

    out = src.with_name(f"{src.stem}_美化版{src.suffix}")
    doc.save(str(out))
    return out


def main():
    for src in TARGETS:
        if not src.exists():
            print(f"[MISS] {src}")
            continue
        out = beautify_doc(src)
        print(f"[DONE] {out}")


if __name__ == "__main__":
    main()

