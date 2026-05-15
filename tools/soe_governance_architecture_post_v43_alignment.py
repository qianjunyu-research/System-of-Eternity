from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


BASE_DOC = Path(
    r"C:\Users\M7120\Documents\New project 7\00_ACTIVE_HANDOFFS\SOE_Upgrade_Waiting_Room\SOE_Architecture_v2_Integration_Snapshot_v0_5_UPGRADED.docx"
)
OUT_DOC = Path(
    r"C:\Users\M7120\Documents\New project 7\00_ACTIVE_HANDOFFS\SOE_Upgrade_Waiting_Room\SOE_Architecture_v2_Integration_Snapshot_v0_5_1_POST_V4_3_ALIGNMENT_CANDIDATE.docx"
)
DESKTOP_FULL = Path(
    r"C:\Users\M7120\Desktop\System of Eternity\SOE Governance Architecture V2\SOE Governance Architecture v0.5 Full Snapshot Upgrade\SOE_Architecture_v2_Integration_Snapshot_v0_5_1_POST_V4_3_ALIGNMENT_CANDIDATE.docx"
)
UPLOAD_DIR = Path(r"C:\Users\M7120\Desktop\SOE_Governance_Architecture_Post_V4_3_ALIGNMENT_UPLOAD_TO_DRIVE")


HERO_FALSIFIABILITY = (
    "Post-V4.3 falsifiability constraint: H(t)/Hero activation requires evidence that the system retains "
    "a recovery pathway independent of Hero intervention. If no such pathway exists, H(t)/Hero activation "
    "is blocked and Trigger B external intervention is the required response instead. This prevents H(t) "
    "from becoming structurally necessary and therefore unchallengeable under Meta Rule 1."
)


POST_V43_NOTES = [
    "P. Post-V4.3 Governance Architecture Alignment Notes",
    (
        "This section records targeted backflow from the completed SOE V4.3 framework line into the "
        "Governance Architecture snapshot line. It does not merge the full V4.3 framework into Governance "
        "Architecture and does not change the evidence boundary of v0.5."
    ),
    (
        "C09 alignment: C09-G05 now treats star-adjacent topology as a recognition blocker rather than only "
        "a scaling blocker. Hub redundancy is allowed only as a documented transitional configuration with a "
        "mandatory reclassification timeline; mesh or federation topology is the required destination. C09-G06 "
        "requires independent third-party attestation that no single actor controls all recognition gates."
    ),
    (
        "Drift alignment: Drift is not a standalone governance component. Drift monitoring is a C07 "
        "Meta-Governance sub-function for long-term integrity monitoring. It detects and routes drift, but "
        "must not become a hidden veto authority."
    ),
    (
        "Hero/H(t) alignment: H(t) remains time-limited and non-routine. Post-V4.3, H(t)/Hero activation also "
        "requires a falsifiability gate: the system must retain a recovery pathway independent of Hero "
        "intervention, or H(t)/Hero activation is blocked and Trigger B external intervention is required."
    ),
    (
        "Chapter 10 alignment note: the Governance Architecture snapshot does not contain the full V4.3 "
        "Chapter 10 crisis-mechanism body. No Chapter 10 text is inserted here. Emergency or crisis-like "
        "mechanisms already present in the architecture line, including H(t), Trigger A, and Trigger B, "
        "inherit C07 trigger-authority constraints."
    ),
    (
        "Evidence boundary: Grand Simulation v0.7 remains bounded simulation evidence for an operationalized "
        "C00-C08 model only. Node v0.1 remains bounded container evidence only. These post-V4.3 alignment "
        "patches do not create simulation evidence, pilot readiness, deployment readiness, empirical "
        "validation, or real-world safety evidence."
    ),
]


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    new_para.add_run(text)
    return new_para


def main() -> None:
    if not BASE_DOC.exists():
        raise FileNotFoundError(BASE_DOC)
    shutil.copy2(BASE_DOC, OUT_DOC)
    doc = Document(OUT_DOC)

    for para in doc.paragraphs:
        if para.text.startswith("Source body: SOE Architecture v2 Integration Snapshot."):
            para.text = (
                "Source body: SOE Architecture v2 Integration Snapshot. Source upgrades: Grand Simulation v0.7 "
                "evidence package, Node v0.1 packet, C00 registry update, C07 v2.3 Meta-Governance patch, "
                "C09 proposal, Node Measurement Protocol, Drift Layer requirements, Psi_extended adversarial "
                "test specification, multi-AI review consolidation, v0.5 blocker list, and post-V4.3 targeted "
                "alignment notes."
            )
            break

    # Hero/H(t) falsifiability alignment.
    for para in doc.paragraphs:
        if para.text.startswith("H(t) is the intervention signal"):
            insert_after(para, HERO_FALSIFIABILITY)
            break

    for table in doc.tables:
        first = table.rows[0].cells[0].text if table.rows and table.rows[0].cells else ""
        if first.startswith("OPEN — H(t) activation criteria"):
            table.rows[0].cells[0].text = (
                "OPEN — H(t) activation criteria, duration limits, trigger authority, abuse prevention mechanisms, "
                "and falsifiability-gate procedures remain high-priority institutional design questions. Post-V4.3, "
                "H(t)/Hero activation must be blocked unless a non-Hero recovery pathway still exists; otherwise "
                "Trigger B external intervention is required."
            )
        if len(table.rows) > 0 and [c.text for c in table.rows[0].cells][:2] == ["Property", "Description"]:
            if any("H(t) is time-limited" in row.cells[0].text for row in table.rows):
                if not any("H(t) requires independent recovery pathway" in row.cells[0].text for row in table.rows):
                    row = table.add_row()
                    row.cells[0].text = "H(t) requires independent recovery pathway"
                    row.cells[1].text = (
                        "Post-V4.3 falsifiability gate: H(t)/Hero activation requires evidence that recovery is "
                        "not structurally dependent on H(t). If no independent recovery path exists, H(t) is blocked "
                        "and Trigger B external intervention is required."
                    )
        if len(table.rows) > 0 and [c.text for c in table.rows[0].cells][:3] == ["Open Item", "Detail", "Priority"]:
            for row in table.rows:
                if row.cells[0].text == "H(t) activation criteria":
                    row.cells[1].text = (
                        "When does H(t) activate? What triggers it? What are the time limits? What is the trigger "
                        "authority structure? Post-V4.3, criteria must include a falsifiability gate proving the "
                        "system retains a non-Hero recovery pathway."
                    )

    # C09 alignment.
    replacements = {
        "C09-G05 Topology fitness: topology classified before recognition; star-adjacent structures redesigned before scaling; federation candidates support cluster monitoring.": (
            "C09-G05 Topology fitness: topology classified before recognition; star-adjacent structures where one participant controls all connections must be redesigned before recognition is granted. Hub redundancy is not a destination topology; it may be used only as a documented transitional configuration with a mandatory reclassification timeline. Mesh or federation topology is the required destination, and Stage 2 recognition requires that the final topology is not star-adjacent."
        ),
        "C09-G06 Anti-capture: no founder, hub, funder, AI operator, security function, or evaluator controls all recognition criteria.": (
            "C09-G06 Anti-capture: no founder, hub, funder, AI operator, security function, or evaluator controls all recognition criteria. Gate satisfaction requires verification by at least one independent party not involved in node formation; self-declaration alone is insufficient, and written attestation is logged in the recognition record."
        ),
    }
    for para in doc.paragraphs:
        if para.text in replacements:
            para.text = replacements[para.text]

    # Drift alignment.
    for para in doc.paragraphs:
        if para.text == "J. Drift Layer Requirements":
            para.text = "J. C07 Long-Term Integrity Monitoring / Drift Requirements"
        elif para.text.startswith("Status: proposed cross-cutting architecture layer only."):
            para.text = (
                "Status: proposed Meta-Governance sub-function only. Drift monitoring is not a standalone governance "
                "component. It is not simulation-validated, not operationally validated, not pilot-ready, and not "
                "deployment-ready."
            )
        elif para.text.startswith("The Drift Layer detects slow deviation"):
            para.text = (
                "C07 long-term integrity monitoring detects slow deviation from SOE ontology, evidence boundaries, "
                "source hierarchy, topology assumptions, detector settings, node-recognition gates, exit rights, "
                "resource dependencies, and archive integrity."
            )
        elif para.text == "The Drift Layer must not become a hidden veto authority.":
            para.text = (
                "Drift monitoring must not become a hidden veto authority. The five-function authority separation "
                "defined for C07 applies: no single actor may perform sensing, classification, routing, response, "
                "and audit for the same drift finding."
            )
        elif para.text == "Drift Layer not simulated or operationalized.":
            para.text = "C07 drift monitoring not simulated or operationalized."
        elif para.text == "Drift Layer v0.2 simulation specification.":
            para.text = "C07 drift monitoring simulation specification."
        elif "v0.7 does not validate C09, Drift Layer" in para.text:
            para.text = para.text.replace("Drift Layer", "C07 drift monitoring")

    # End notes.
    if not any(p.text == POST_V43_NOTES[0] for p in doc.paragraphs):
        anchor = None
        for para in reversed(doc.paragraphs):
            if para.text.startswith("A future v0.8 simulation should"):
                anchor = para
                break
        if anchor is None:
            anchor = doc.paragraphs[-1]
        current = anchor
        for note in POST_V43_NOTES:
            current = insert_after(current, note)

    doc.save(OUT_DOC)
    DESKTOP_FULL.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_DOC, DESKTOP_FULL)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_DOC, UPLOAD_DIR / OUT_DOC.name)
    print(f"candidate={OUT_DOC}")
    print(f"desktop_copy={DESKTOP_FULL}")
    print(f"upload_copy={UPLOAD_DIR / OUT_DOC.name}")


if __name__ == "__main__":
    main()
