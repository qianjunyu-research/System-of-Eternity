from __future__ import annotations

import argparse
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


W_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def extract_paragraphs(path: Path) -> list[str]:
    with ZipFile(path) as archive:
        data = archive.read("word/document.xml")
    root = ET.fromstring(data)
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", W_NS):
        texts = [node.text for node in paragraph.findall(".//w:t", W_NS) if node.text]
        text = "".join(texts).strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True


def add_header(doc: Document, title_text: str) -> None:
    header = doc.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(title_text)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)


def add_title_block(doc: Document, title_text: str) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(title_text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Formatted Research Model Draft")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_summary_table(doc: Document, paragraphs: list[str]) -> None:
    scope_items = [p[2:].strip() for p in paragraphs if p.startswith("- ")]
    state_vars = [p for p in paragraphs if re.match(r"^[TDCS]\(t\):", p)]
    classification = next((p for p in paragraphs if p.startswith("SOE v1.3 is ")), "")

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    labels = ("Scope", "State Variables", "Classification")
    values = (
        "; ".join(scope_items[:5]) if scope_items else "",
        "; ".join(state_vars),
        classification,
    )

    for row_index, (label, value) in enumerate(zip(labels, values)):
        label_cell = table.cell(row_index, 0)
        value_cell = table.cell(row_index, 1)
        label_cell.text = label
        value_cell.text = value
        set_cell_shading(label_cell, "D9EAF7")

    doc.add_paragraph()


def add_labeled_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    label, content = text.split(":", 1)
    label_run = paragraph.add_run(f"{label.strip()}: ")
    label_run.bold = True
    value_run = paragraph.add_run(content.strip())
    label_run.font.name = value_run.font.name = "Times New Roman"
    label_run.font.size = value_run.font.size = Pt(11)


def add_body(doc: Document, paragraphs: list[str]) -> None:
    numbered_heading = re.compile(r"^\d+\.\s+")
    sub_heading = re.compile(r"^\d+\.\d+\s+")

    for text in paragraphs[1:]:
        if numbered_heading.match(text):
            doc.add_paragraph(text, style="Heading 1")
            continue
        if sub_heading.match(text):
            doc.add_paragraph(text, style="Heading 2")
            continue
        if text == "Definition:" or text == "Condition:" or text == "Behavior:" or text == "Interpretation:":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(text)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            continue
        if text.startswith("- "):
            doc.add_paragraph(text[2:].strip(), style="List Bullet")
            continue
        if re.match(r"^(System Name|Primary variable|Secondary variable|Tertiary variable|Dimension \d+|Type [A-D]|Property \d+)", text):
            doc.add_paragraph(text)
            continue
        if re.match(r"^[A-Z][A-Za-z ]+:", text) or re.match(r"^[TDCS]\(t\):", text):
            add_labeled_paragraph(doc, text)
            continue
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(6)


def format_document(source: Path, destination: Path) -> None:
    paragraphs = extract_paragraphs(source)
    if not paragraphs:
        raise ValueError(f"No paragraphs found in {source}")

    doc = Document()
    configure_document(doc)
    add_header(doc, "SOE Research Model")
    add_title_block(doc, paragraphs[0])
    add_summary_table(doc, paragraphs)
    add_body(doc, paragraphs)
    doc.save(destination)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Format a research DOCX into a cleaner academic-style layout.")
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    format_document(args.source, args.destination)
    print(f"Saved formatted document to {args.destination}")


if __name__ == "__main__":
    main()
