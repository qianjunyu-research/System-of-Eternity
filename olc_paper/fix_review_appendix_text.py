from pathlib import Path
from docx import Document


TARGET = Path(r"C:\Users\sosoy\Downloads\review_with_original_full.docx")
TARGET_CN = Path(r"C:\Users\sosoy\Downloads\终极了悟_结构点评_含原文完整附录.docx")


def fix_doc(path: Path) -> None:
    if not path.exists():
        return
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "???????????":
            p.text = "附录：原文（完整复刻）"
        elif t == "???????????????????????????":
            p.text = "以下内容为原文完整附录，保留原始结构与图示（如存在）。"
        elif (
            t == "http://www.fodizi.net/qt/dazhaofashi/21185.html"
            or "fodizi.net/qt/dazhaofashi/21185.html" in t
        ):
            p.text = "参考链接（Fodizi）：https://www.fodizi.net/qt/dazhaofashi/21185.html"
    doc.save(str(path))


def main() -> None:
    fix_doc(TARGET)
    # Keep a Chinese-filename copy in sync
    Document(str(TARGET)).save(str(TARGET_CN))
    print(TARGET)
    print(TARGET_CN)


if __name__ == "__main__":
    main()

