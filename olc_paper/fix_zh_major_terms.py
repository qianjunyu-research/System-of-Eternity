from pathlib import Path
from docx import Document


SRC = Path(r"C:\Users\sosoy\Downloads\OLC_Paper_Formatted_zhCN.docx")
DST = SRC.with_name("OLC_Paper_Formatted_zhCN_majorfix_FINAL.docx")


REPLACEMENTS = {
    "反向-负载-耦合": "对立-负载-耦合",
    "抽象的": "摘要",
    "一、简介": "一、引言",
    "预测处理帐户": "预测处理理论",
    "紧急自主": "涌现自主性",
    "涌现自治": "涌现自主性",
    "智能体通过故意脱离而被拆除": "该结构通过有意识的脱离而被溶解",
    "内部代理动态": "内部主体动态",
    "不适合复制": "不应复制",
    "代理式结构": "智能体式结构",
    "内部认知代理在这里被定义为": "内部认知主体在这里被定义为",
    "可溶解性": "可消解性",
    "权利要求仅限于": "本文论断仅限于",
    "4.反对-负载-耦合（OLC）模型": "4. 对立-负载-耦合（OLC）模型",
    "反对-负载-耦合": "对立-负载-耦合",
    "共享值引用": "共享价值锚点",
    "客观空间": "目标空间",
    "关键的反面是": "关键反转点在于",
    "反对压力": "对立压力",
    "反对派渠道": "对立通道",
    "最终解体": "最终溶解",
    "解散是稳定的": "溶解过程带来了稳定性",
    "代理后过渡": "后主体转型",
    "解体后": "溶解后",
    "基于主体的内部处理": "基于智能体的内部处理",
    "边界帐户": "边界说明",
    "患病率": "发生率",
    "多主体验证": "多受试者验证",
    "2026?": "2026年",
}


def replace_text(text: str) -> str:
    out = text
    for old, new in REPLACEMENTS.items():
        out = out.replace(old, new)
    return out


def process_paragraphs(paragraphs):
    for p in paragraphs:
        new = replace_text(p.text)
        if new != p.text:
            p.text = new


def main():
    doc = Document(str(SRC))

    process_paragraphs(doc.paragraphs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    for sec in doc.sections:
        process_paragraphs(sec.header.paragraphs)
        process_paragraphs(sec.footer.paragraphs)

    doc.save(str(DST))
    print(DST)


if __name__ == "__main__":
    main()

