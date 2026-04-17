from pathlib import Path
from docx import Document


SRC = Path(r"C:\Users\sosoy\Downloads\OLC_Paper_Formatted_zhCN_majorfix_FINAL.docx")
DST = SRC.with_name("OLC_Paper_Formatted_zhCN_majorfix_FINAL_v2.docx")


REPLACEMENTS = {
    "、解体和": "、溶解和",
    "解体后": "溶解后",
    "通过系统的脱离可以实现解散": "通过系统的脱离可以实现溶解",
}


def apply_replacements(text: str) -> str:
    out = text
    for old, new in REPLACEMENTS.items():
        out = out.replace(old, new)
    return out


def process_paragraphs(paragraphs):
    for p in paragraphs:
        text = p.text
        new_text = apply_replacements(text)
        if text.strip() == "2026":
            new_text = "2026年"
        if new_text != text:
            p.text = new_text


def main():
    doc = Document(str(SRC))

    process_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    for sec in doc.sections:
        process_paragraphs(sec.header.paragraphs)
        process_paragraphs(sec.footer.paragraphs)

    doc.save(str(DST))
    print(DST)


if __name__ == "__main__":
    main()

