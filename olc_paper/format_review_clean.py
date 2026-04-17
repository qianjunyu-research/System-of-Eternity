from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


SRC = Path(r"C:\Users\sosoy\Downloads\review_target.docx")
OUT_ASCII = Path(r"C:\Users\sosoy\Downloads\review_formatted_clean.docx")
OUT_CN = Path(r"C:\Users\sosoy\Downloads\终极了悟_结构点评_格式化版_无乱码.docx")


def main() -> None:
    src_doc = Document(str(SRC))
    paras = [p.text.strip() for p in src_doc.paragraphs if p.text.strip()]

    start = None
    for i, t in enumerate(paras):
        if "结构点评：关于“终极觉悟与修行路径”的结构分析" in t:
            start = i
            break
    if start is None:
        for i, t in enumerate(paras):
            if t.startswith("1. 原文核心观点"):
                start = max(0, i - 2)
                break
    if start is None:
        start = 0

    content = paras[start:]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)

    title = doc.add_paragraph("结构点评（格式化版）")
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("主题：关于“终极觉悟与修行路径”的结构分析")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if sub.runs:
        sub.runs[0].bold = True

    doc.add_paragraph("")

    num_heading = re.compile(r"^\d+\.\s*")
    sub_heading = re.compile(r"^\d+\.\d+\s*")

    for t in content:
        if t == "📄":
            continue

        if num_heading.match(t) and not sub_heading.match(t):
            doc.add_heading(t, level=1)
            continue

        if sub_heading.match(t):
            doc.add_heading(t, level=2)
            continue

        if t.startswith("👉") or t.startswith("❗"):
            p = doc.add_paragraph(t, style="List Bullet")
            if p.runs:
                p.runs[0].bold = True
            continue

        if t.startswith("（1）") or t.startswith("（2）"):
            p = doc.add_paragraph(t)
            if p.runs:
                p.runs[0].bold = True
            continue

        if t.startswith("A.") or t.startswith("B.") or t.startswith("C."):
            doc.add_paragraph(t, style="List Bullet")
            continue

        if "↓" in t or "→" in t:
            p = doc.add_paragraph(t)
            for r in p.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(10.5)
            continue

        doc.add_paragraph(t)

    doc.add_paragraph("")
    foot = doc.add_paragraph("注：本稿仅做结构与表达格式化，不改变原观点结论。")
    if foot.runs:
        foot.runs[0].italic = True

    doc.save(str(OUT_ASCII))
    # Also save a Chinese filename version
    doc.save(str(OUT_CN))
    print(OUT_ASCII)
    print(OUT_CN)


if __name__ == "__main__":
    main()

